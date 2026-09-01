from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Times New Roman", size=10, color="222222")
THIN_GRAY = Side(style="thin", color="D9DEE3")


def _format_workbook(writer: pd.ExcelWriter) -> None:
    """Apply compact report-table formatting to every populated worksheet."""
    workbook = writer.book
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = (
            sheet.ORIENTATION_LANDSCAPE if sheet.max_column > 5 else sheet.ORIENTATION_PORTRAIT
        )
        sheet.page_margins.left = 0.35
        sheet.page_margins.right = 0.35
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.print_options.horizontalCentered = True
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN_GRAY)
        sheet.row_dimensions[1].height = 24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = BODY_FONT
                cell.alignment = Alignment(
                    horizontal="right" if isinstance(cell.value, (int, float)) else "left",
                    vertical="center",
                )
        for column_cells in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(max(max((len(value) for value in values), default=0) + 2, 11), 28)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = sheet.cell(1, cell.column).value or ""
                header = str(header)
                if not isinstance(cell.value, (int, float)):
                    continue
                if header in {
                    "n_common_years", "ASK_rank", "RPK_rank", "n_years",
                    "有效年份数", "观测数", "数值",
                }:
                    cell.number_format = "0"
                elif header in {"R", "ln_R"}:
                    cell.number_format = "0.000"
                elif any(key in header for key in ["PLF", "比例"]):
                    cell.number_format = "0.0"
                elif any(key in header for key in ["mean_ASK", "mean_RPK", "sum_ASK", "sum_RPK"]):
                    cell.number_format = "#,##0.00"


