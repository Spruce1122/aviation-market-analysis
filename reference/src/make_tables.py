"""Generate T01–T04 from shared processed DataFrames."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd

from .config import (
    DIRECTION_RANK_TOP_N,
    RANKING_END_YEAR,
    RANKING_MIN_COMMON_YEARS,
    RANKING_START_YEAR,
    RANKING_TOP_N,
    TABLE_DIR,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Times New Roman", size=10, color="222222")
THIN_GRAY = Side(style="thin", color="D9DEE3")


def _format_workbook(writer: pd.ExcelWriter) -> None:
    """Apply compact report-table formatting to every populated worksheet."""
    workbook = writer.book
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = (
            sheet.ORIENTATION_LANDSCAPE if sheet.max_column > 5 else sheet.ORIENTATION_PORTRAIT
        )
        sheet.page_margins.left = 0.35
        sheet.page_margins.right = 0.35
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.print_options.horizontalCentered = True
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=THIN_GRAY)
        sheet.row_dimensions[1].height = 24
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = BODY_FONT
                cell.alignment = Alignment(
                    horizontal="right" if isinstance(cell.value, (int, float)) else "left",
                    vertical="center",
                )
        for column_cells in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(max(max((len(value) for value in values), default=0) + 2, 11), 28)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = sheet.cell(1, cell.column).value or ""
                header = str(header)
                if not isinstance(cell.value, (int, float)):
                    continue
                if header in {
                    "n_common_years", "ASK_rank", "RPK_rank", "n_years",
                    "有效年份数", "观测数", "数值",
                }:
                    cell.number_format = "0"
                elif header in {"R", "ln_R"}:
                    cell.number_format = "0.000"
                elif any(key in header for key in ["PLF", "比例"]):
                    cell.number_format = "0.0"
                elif any(key in header for key in ["mean_ASK", "mean_RPK", "sum_ASK", "sum_RPK"]):
                    cell.number_format = "#,##0.00"


def _write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_workbook(writer)
    return path


def make_t01(metrics: pd.DataFrame, logger: logging.Logger) -> tuple[Path, pd.DataFrame]:
    sample = metrics.loc[
        metrics["Time"].between(RANKING_START_YEAR, RANKING_END_YEAR)
        & metrics["ASKs"].gt(0)
        & metrics["RPKs"].gt(0)
    ].copy()
    ranking = (
        sample.groupby(["Country Name", "Country Code"], as_index=False)
        .agg(
            n_common_years=("Time", "nunique"),
            mean_ASK=("ASKs", "mean"),
            mean_RPK=("RPKs", "mean"),
            sum_ASK=("ASKs", "sum"),
            sum_RPK=("RPKs", "sum"),
        )
    )
    ranking = ranking.loc[ranking["n_common_years"].ge(RANKING_MIN_COMMON_YEARS)].copy()
    ranking["ASK_rank"] = ranking["mean_ASK"].rank(method="min", ascending=False).astype(int)
    ranking["RPK_rank"] = ranking["mean_RPK"].rank(method="min", ascending=False).astype(int)
    ranking["weighted_PLF"] = ranking["sum_RPK"].div(ranking["sum_ASK"]).mul(100)
    output_columns = [
        "Country Name", "Country Code", "n_common_years", "mean_ASK", "mean_RPK",
        "ASK_rank", "RPK_rank", "weighted_PLF",
    ]
    ask_top = ranking.sort_values(["ASK_rank", "Country Code"])[output_columns].head(RANKING_TOP_N)
    rpk_top = ranking.sort_values(["RPK_rank", "Country Code"])[output_columns].head(RANKING_TOP_N)

    path = _write_excel(
        TABLE_DIR / "T01_主要航空市场ASK_RPK排名.xlsx",
        {"ASK_Top10": ask_top, "RPK_Top10": rpk_top},
    )
    ask_top.to_csv(TABLE_DIR / "T01_ASK_Top10.csv", index=False, encoding="utf-8-sig")
    rpk_top.to_csv(TABLE_DIR / "T01_RPK_Top10.csv", index=False, encoding="utf-8-sig")
    logger.info("T01完成，合格国家数=%d", len(ranking))
    return path, ranking


def make_t02(metrics: pd.DataFrame, logger: logging.Logger) -> Path:
    common = metrics.loc[metrics["common_growth_sample"]].copy()
    total = len(common)
    categories = [
        ("同时增长", "both_up"),
        ("同时下降", "both_down"),
        ("ASK增长、RPK下降", "ASK_up_RPK_down"),
        ("ASK下降、RPK增长", "ASK_down_RPK_up"),
        ("同方向合计", "same_direction"),
        ("反方向合计", "opposite_direction"),
    ]
    rows = []
    for label, column in categories:
        count = int(common[column].sum())
        rows.append(
            {"变化方向": label, "观测数": count, "占有效样本比例（%）": count / total * 100}
        )
    summary = pd.DataFrame(rows)
    check = pd.DataFrame(
        [
            {"校验项目": "共同增长有效样本总数", "数值": total},
            {"校验项目": "增长率等于0且未进入严格增减类别", "数值": int(common["zero_growth_unclassified"].sum())},
            {"校验项目": "四个互斥基本类别合计", "数值": int(
                common[["both_up", "both_down", "ASK_up_RPK_down", "ASK_down_RPK_up"]].sum().sum()
            )},
        ]
    )
    path = _write_excel(
        TABLE_DIR / "T02_ASK_RPK年度变化方向统计.xlsx",
        {"方向统计": summary, "校验": check},
    )
    logger.info("T02完成，共同增长样本=%d", total)
    return path


def make_t03(direction_low: pd.DataFrame, direction_high: pd.DataFrame, logger: logging.Logger) -> Path:
    columns = ["country_name", "country_code", "mean_ASK_out", "mean_ASK_in", "R", "ln_R", "n_years"]
    low = direction_low[columns].sort_values(["ln_R", "country_code"])
    high = direction_high[columns].sort_values(["ln_R", "country_code"], ascending=[False, True])
    path = _write_excel(
        TABLE_DIR / "T03_ASK_out_in不对称国家.xlsx",
        {"ln_R最低9国": low, "ln_R最高9国": high},
    )
    logger.info("T03完成")
    return path


def make_t04(direction_metrics: pd.DataFrame, logger: logging.Logger) -> Path:
    columns = ["country_name", "country_code", "mean_ASK_out", "mean_ASK_in", "R", "ln_R", "n_years"]
    rename = {
        "country_name": "国家",
        "country_code": "ISO",
        "n_years": "有效年份数",
    }
    out_top = (
        direction_metrics.nlargest(DIRECTION_RANK_TOP_N, ["mean_ASK_out", "mean_ASK_in"])[columns]
        .rename(columns=rename)
    )
    in_top = (
        direction_metrics.nlargest(DIRECTION_RANK_TOP_N, ["mean_ASK_in", "mean_ASK_out"])[columns]
        .rename(columns=rename)
    )
    path = _write_excel(
        TABLE_DIR / "T04_ASK_out_in_Top10.xlsx",
        {"ASK_out_top10": out_top, "ASK_in_top10": in_top},
    )
    logger.info("T04完成")
    return path


def make_all_tables(
    metrics: pd.DataFrame,
    direction_metrics: pd.DataFrame,
    direction_low: pd.DataFrame,
    direction_high: pd.DataFrame,
    logger: logging.Logger,
) -> tuple[list[Path], pd.DataFrame]:
    t01, ranking = make_t01(metrics, logger)
    return [
        t01,
        make_t02(metrics, logger),
        make_t03(direction_low, direction_high, logger),
        make_t04(direction_metrics, logger),
    ], ranking
